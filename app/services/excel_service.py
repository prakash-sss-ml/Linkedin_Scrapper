from __future__ import annotations

from datetime import date
from difflib import SequenceMatcher
from pathlib import Path
import re
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import (
    get_column_letter,
)
from openpyxl.worksheet.table import (
    Table,
    TableStyleInfo,
)

from app.models import (
    ExtractionError,
    JobRecord,
)
from app.services.apollo_service import (
    ApolloEnrichmentResult,
)


class ExcelService:
    """
    Append all extracted jobs into one persistent Excel workbook.

    Every extracted job is written regardless of Apollo match status.
    If a job has multiple Apollo officials, one row is written per official.
    If Apollo has no official or no company match, one job row is still written
    with the available extracted/Apollo fields and blank official columns.
    """

    HEADERS = [
        "Job ID",
        "Job Title",
        "Extracted Company Name",
        "Location",
        "Workplace Type",
        "Post Contact Name",
        "Post Contact Email ID",
        "Job Description",

        "Apollo Match Status",
        "Apollo Company Name",
        "Company Website",
        "Industry",
        "Employee Count",
        "Annual Revenue",
        "Company City",
        "Company Country",

        "Official Name",
        "Official Title",
        "Official Work Email",
        "Official LinkedIn",

        "Apollo Enrichment Error",
        "Date",
    ]

    def create_workbook(
        self,
        jobs: list[JobRecord],
        errors: list[ExtractionError],
        output_path: Path,
        apollo_results: dict[
            str,
            ApolloEnrichmentResult,
        ] | None = None,
    ) -> Path:
        del errors

        apollo_results = apollo_results or {}
        output_path = Path(output_path)

        # Use exactly the path supplied by main.py. main.py owns the decision
        # that this is one persistent workbook shared by all runs.
        output_path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        if output_path.exists():
            workbook = load_workbook(output_path)

            if "Job_Postings" in workbook.sheetnames:
                worksheet = workbook["Job_Postings"]
            else:
                worksheet = workbook.create_sheet(
                    "Job_Postings"
                )
                self._write_headers(worksheet)

            # Upgrade older workbooks that may be missing the post-contact email
            # column and/or the Date column.
            self._ensure_schema_columns(worksheet)
            self._validate_headers(worksheet)

            # Remove duplicate rows that may already exist from repeated
            # LinkedIn posts captured in earlier recordings.
            self._remove_duplicate_rows(worksheet)
        else:
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Job_Postings"
            self._write_headers(worksheet)

        row_index = max(2, worksheet.max_row + 1)

        # Existing row keys prevent the exact same official row being written
        # twice. Existing post keys prevent a LinkedIn post captured again in a
        # later recording from adding a second set of Apollo officials.
        existing_row_keys = self._existing_row_keys(worksheet)
        existing_post_keys = self._existing_post_keys(worksheet)

        for job in jobs:
            post_key = self._post_key_from_job(job)
            if self._contains_post_key(post_key, existing_post_keys):
                continue

            enrichment = apollo_results.get(
                job.job_id
            )

            # Write every extracted job, regardless of Apollo match status.
            # Apollo enrichment may be missing, NO_MATCH, API_ERROR, etc.
            officials = enrichment.officials if enrichment is not None else []

            # One separate Excel row for every matched official. A repeated
            # recording of this post is skipped at the job level above, so a
            # later Apollo result cannot add another copy of the same post.
            if officials:
                for official in officials:
                    values = self._job_values(
                        job=job,
                        enrichment=enrichment,
                        official=official,
                    )

                    row_key = self._dedupe_key_from_values(values)
                    if row_key in existing_row_keys:
                        continue

                    self._write_row(
                        worksheet=worksheet,
                        row_index=row_index,
                        values=values,
                        job=job,
                    )

                    existing_row_keys.add(row_key)
                    row_index += 1
            else:
                # Keep one row even when Apollo did not return a company match
                # or a contactable official.
                values = self._job_values(
                    job=job,
                    enrichment=enrichment,
                    official=None,
                )

                row_key = self._dedupe_key_from_values(values)
                if row_key not in existing_row_keys:
                    self._write_row(
                        worksheet=worksheet,
                        row_index=row_index,
                        values=values,
                        job=job,
                    )
                    existing_row_keys.add(row_key)
                    row_index += 1

            existing_post_keys.append(post_key)

        self._format_sheet(worksheet)
        self._sync_table(worksheet)

        workbook.save(output_path)

        return output_path

    @staticmethod
    def _is_matched(
        enrichment: ApolloEnrichmentResult | None,
    ) -> bool:
        if enrichment is None:
            return False

        return (
            str(enrichment.company.match_status or "")
            .strip()
            .casefold()
            == "matched"
        )

    def _ensure_schema_columns(
        self,
        worksheet: Any,
    ) -> None:
        """Upgrade supported older workbooks to the current Excel schema.

        Supported legacy layouts include:
        - the current exporter before ``Post Contact Email ID`` was restored;
        - the older exporter that used the header ``Post Contact Email``;
        - either of those layouts before the ``Date`` column was added.

        When the email column is missing, insert it immediately after
        ``Post Contact Name`` so existing row data stays aligned.
        """
        existing_headers = [
            worksheet.cell(row=1, column=index).value
            for index in range(1, worksheet.max_column + 1)
        ]

        while existing_headers and existing_headers[-1] in (None, ""):
            existing_headers.pop()

        if existing_headers == self.HEADERS:
            return

        headers_without_email = [
            header
            for header in self.HEADERS
            if header != "Post Contact Email ID"
        ]
        headers_without_email_or_date = headers_without_email[:-1]

        headers_with_old_email_name = [
            (
                "Post Contact Email"
                if header == "Post Contact Email ID"
                else header
            )
            for header in self.HEADERS
        ]
        headers_with_old_email_name_no_date = headers_with_old_email_name[:-1]

        existing_headers_tuple = tuple(existing_headers)

        if existing_headers_tuple in {
            tuple(headers_without_email),
            tuple(headers_without_email_or_date),
        }:
            # Column numbers are 1-based for openpyxl. Insert directly after
            # Post Contact Name; all existing columns/data shift right safely.
            contact_name_column = (
                headers_without_email.index("Post Contact Name") + 1
            )
            worksheet.insert_cols(contact_name_column + 1, 1)
            self._write_headers(worksheet)
            return

        if existing_headers_tuple in {
            tuple(headers_with_old_email_name),
            tuple(headers_with_old_email_name_no_date),
        }:
            # The data column already exists; only normalize the header name
            # and add Date when it was absent.
            self._write_headers(worksheet)
            return

    def _validate_headers(
        self,
        worksheet: Any,
    ) -> None:
        existing_headers = [
            worksheet.cell(
                row=1,
                column=index,
            ).value
            for index in range(
                1,
                len(self.HEADERS) + 1,
            )
        ]

        if existing_headers != self.HEADERS:
            raise RuntimeError(
                "Existing Apollo master Excel has unexpected headers: "
                f"{worksheet.title}"
            )

    def _remove_non_matched_rows(
        self,
        worksheet: Any,
    ) -> None:
        status_column = self._column_index(
            "Apollo Match Status"
        )

        for row_index in range(
            worksheet.max_row,
            1,
            -1,
        ):
            status = str(
                worksheet.cell(
                    row=row_index,
                    column=status_column,
                ).value
                or ""
            ).strip().casefold()

            if status != "matched":
                worksheet.delete_rows(row_index, 1)

    def _remove_duplicate_rows(
        self,
        worksheet: Any,
    ) -> None:
        """Remove duplicate exported rows while keeping the first occurrence."""
        seen: set[tuple[str, ...]] = set()

        # Walk from top to bottom, remember duplicate row numbers, then delete
        # bottom-to-top so row indexes do not shift underneath us.
        duplicate_rows: list[int] = []
        for row_index in range(2, worksheet.max_row + 1):
            values = [
                worksheet.cell(
                    row=row_index,
                    column=column_index,
                ).value
                for column_index in range(1, len(self.HEADERS) + 1)
            ]
            row_key = self._dedupe_key_from_values(values)

            if row_key in seen:
                duplicate_rows.append(row_index)
            else:
                seen.add(row_key)

        for row_index in reversed(duplicate_rows):
            worksheet.delete_rows(row_index, 1)

    def _existing_row_keys(
        self,
        worksheet: Any,
    ) -> set[tuple[str, ...]]:
        """Return stable row keys for all rows already in the workbook."""
        keys: set[tuple[str, ...]] = set()

        for row_index in range(2, worksheet.max_row + 1):
            values = [
                worksheet.cell(
                    row=row_index,
                    column=column_index,
                ).value
                for column_index in range(1, len(self.HEADERS) + 1)
            ]
            keys.add(self._dedupe_key_from_values(values))

        return keys

    def _existing_post_keys(
        self,
        worksheet: Any,
    ) -> list[tuple[str, str, str, str]]:
        """Return LinkedIn-post identities already stored in the workbook.

        Header lookup is name-based so duplicate filtering still works when an
        older workbook does not yet contain the restored post-contact email
        column.
        """
        keys: list[tuple[str, str, str, str]] = []

        header_to_column = {
            str(worksheet.cell(row=1, column=column_index).value or "").strip():
                column_index
            for column_index in range(1, worksheet.max_column + 1)
        }

        required_headers = (
            "Job Title",
            "Extracted Company Name",
            "Location",
            "Workplace Type",
        )
        if any(header not in header_to_column for header in required_headers):
            return keys

        for row_index in range(2, worksheet.max_row + 1):
            key = (
                self._normalize_dedupe_value(
                    worksheet.cell(
                        row=row_index,
                        column=header_to_column["Job Title"],
                    ).value
                ),
                self._normalize_dedupe_value(
                    worksheet.cell(
                        row=row_index,
                        column=header_to_column["Extracted Company Name"],
                    ).value
                ),
                self._normalize_dedupe_value(
                    worksheet.cell(
                        row=row_index,
                        column=header_to_column["Location"],
                    ).value
                ),
                self._normalize_dedupe_value(
                    worksheet.cell(
                        row=row_index,
                        column=header_to_column["Workplace Type"],
                    ).value
                ),
            )
            if not self._contains_post_key(key, keys):
                keys.append(key)

        return keys

    def filter_new_jobs(
        self,
        jobs: list[JobRecord],
        output_path: Path,
    ) -> list[JobRecord]:
        """
        Remove repeated LinkedIn posts before Apollo enrichment.

        This checks both duplicates inside the current extraction and posts that
        are already present in the persistent Excel workbook. This saves Apollo
        calls as well as preventing duplicate Excel rows.
        """
        output_path = Path(output_path)
        existing_post_keys: list[tuple[str, str, str, str]] = []

        if output_path.exists():
            workbook = load_workbook(
                output_path,
                read_only=True,
                data_only=False,
            )
            try:
                if "Job_Postings" in workbook.sheetnames:
                    worksheet = workbook["Job_Postings"]
                    header_values = [
                        worksheet.cell(row=1, column=index).value
                        for index in range(1, worksheet.max_column + 1)
                    ]
                    required_post_headers = {
                        "Job Title",
                        "Extracted Company Name",
                        "Location",
                        "Workplace Type",
                    }
                    if required_post_headers.issubset(set(header_values)):
                        existing_post_keys = self._existing_post_keys(worksheet)
            finally:
                workbook.close()

        filtered: list[JobRecord] = []
        seen_keys = list(existing_post_keys)

        for job in jobs:
            key = self._post_key_from_job(job)
            if self._contains_post_key(key, seen_keys):
                continue
            seen_keys.append(key)
            filtered.append(job)

        return filtered

    def _dedupe_key_from_values(
        self,
        values: list[Any],
    ) -> tuple[str, ...]:
        """Identify one exported contact row without Job ID, Date or OCR text."""
        value_by_header = {
            header: (values[index] if index < len(values) else "")
            for index, header in enumerate(self.HEADERS)
        }

        post_key = self._post_key_from_values(values)

        # Use a single stable official identity. Do not include title, job
        # description or post-contact OCR because those can change between
        # captures even when it is the same LinkedIn post/person.
        official_identity = (
            self._normalize_dedupe_value(value_by_header["Official Name"])
            or self._normalize_dedupe_value(value_by_header["Official LinkedIn"])
            or self._normalize_dedupe_value(value_by_header["Official Work Email"])
        )

        return (*post_key, official_identity)

    def _post_key_from_values(
        self,
        values: list[Any],
    ) -> tuple[str, str, str, str]:
        value_by_header = {
            header: (values[index] if index < len(values) else "")
            for index, header in enumerate(self.HEADERS)
        }

        return (
            self._normalize_dedupe_value(value_by_header["Job Title"]),
            self._normalize_dedupe_value(value_by_header["Extracted Company Name"]),
            self._normalize_dedupe_value(value_by_header["Location"]),
            self._normalize_dedupe_value(value_by_header["Workplace Type"]),
        )

    def _post_key_from_job(
        self,
        job: JobRecord,
    ) -> tuple[str, str, str, str]:
        return (
            self._normalize_dedupe_value(job.job_title),
            self._normalize_dedupe_value(job.company_name),
            self._normalize_dedupe_value(job.location),
            self._normalize_dedupe_value(job.workplace_type),
        )

    def _contains_post_key(
        self,
        candidate: tuple[str, str, str, str],
        existing: list[tuple[str, str, str, str]],
    ) -> bool:
        return any(
            self._post_keys_match(candidate, previous)
            for previous in existing
        )

    @staticmethod
    def _post_keys_match(
        left: tuple[str, str, str, str],
        right: tuple[str, str, str, str],
    ) -> bool:
        left_title, left_company, left_location, left_workplace = left
        right_title, right_company, right_location, right_workplace = right

        if not left_title or not right_title:
            return False
        if not left_company or not right_company:
            return False

        title_match = (
            left_title == right_title
            or SequenceMatcher(None, left_title, right_title).ratio() >= 0.90
        )
        company_match = (
            left_company == right_company
            or SequenceMatcher(None, left_company, right_company).ratio() >= 0.92
        )

        if not title_match or not company_match:
            return False

        # Missing location/workplace data must not make the same post unique.
        location_match = (
            not left_location
            or not right_location
            or left_location == right_location
            or SequenceMatcher(None, left_location, right_location).ratio() >= 0.90
        )
        workplace_match = (
            not left_workplace
            or not right_workplace
            or left_workplace == right_workplace
        )

        return location_match and workplace_match

    @staticmethod
    def _normalize_dedupe_value(value: Any) -> str:
        if value is None:
            return ""

        normalized = " ".join(str(value).split()).casefold()
        normalized = re.sub(r"[^a-z0-9]+", " ", normalized)
        return " ".join(normalized.split())

    def _job_values(
        self,
        job: JobRecord,
        enrichment: (
            ApolloEnrichmentResult
            | None
        ),
        official: Any | None,
    ) -> list[Any]:
        if enrichment is None:
            company = None
            enrichment_error = ""
        else:
            company = enrichment.company
            enrichment_error = enrichment.error

        company_values = [
            (
                company.match_status
                if company
                else "NOT_RUN"
            ),
            (
                company.name
                if company
                else ""
            ),
            (
                company.website_url
                if company
                else ""
            ),
            (
                company.industry
                if company
                else ""
            ),
            (
                company.employee_count
                if company
                else ""
            ),
            (
                company.annual_revenue
                if company
                else ""
            ),
            (
                company.city
                if company
                else ""
            ),
            (
                company.country
                if company
                else ""
            ),
        ]

        if official is None:
            official_values = [
                "",
                "",
                "",
                "",
            ]
        else:
            official_values = [
                official.full_name,
                official.title,
                official.work_email,
                official.linkedin_url,
            ]

        return [
            job.job_id,
            job.job_title,
            job.company_name,
            job.location,
            job.workplace_type,
            job.contact_name,
            self._join_values(
                getattr(job, "contact_emails", None)
            ),
            job.job_description,

            *company_values,
            *official_values,

            enrichment_error,
            date.today(),
        ]

    def _write_headers(
        self,
        worksheet: Any,
    ) -> None:
        header_fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )

        header_font = Font(
            color="FFFFFF",
            bold=True,
            size=11,
        )

        border = self._thin_border()

        for column_index, header in enumerate(
            self.HEADERS,
            start=1,
        ):
            cell = worksheet.cell(
                row=1,
                column=column_index,
                value=header,
            )

            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        worksheet.row_dimensions[1].height = 42

    def _write_row(
        self,
        worksheet: Any,
        row_index: int,
        values: list[Any],
        job: JobRecord,
    ) -> None:
        border = self._thin_border()

        alternate_fill = PatternFill(
            fill_type="solid",
            fgColor="F4F8FC",
        )

        left_aligned_headers = {
            "Job Title",
            "Extracted Company Name",
            "Location",
            "Post Contact Name",
            "Post Contact Email ID",
            "Job Description",
            "Apollo Company Name",
            "Company Website",
            "Industry",
            "Company City",
            "Company Country",
            "Official Name",
            "Official Title",
            "Official Work Email",
            "Official LinkedIn",
            "Apollo Enrichment Error",
        }

        for column_index, value in enumerate(
            values,
            start=1,
        ):
            cell = worksheet.cell(
                row=row_index,
                column=column_index,
                value=self._safe_excel_value(
                    value
                ),
            )

            cell.border = border

            if row_index % 2 == 0:
                cell.fill = alternate_fill

            header = self.HEADERS[
                column_index - 1
            ]

            cell.alignment = Alignment(
                vertical="top",
                horizontal=(
                    "left"
                    if header in left_aligned_headers
                    else "center"
                ),
                wrap_text=True,
            )

        # Extraction date
        worksheet.cell(
            row=row_index,
            column=self._column_index("Date"),
        ).number_format = "yyyy-mm-dd"

        # Annual Revenue
        worksheet.cell(
            row=row_index,
            column=self._column_index(
                "Annual Revenue"
            ),
        ).number_format = "#,##0.00"

        # Apollo match status.
        apollo_cell = worksheet.cell(
            row=row_index,
            column=self._column_index(
                "Apollo Match Status"
            ),
        )

        apollo_status = str(
            apollo_cell.value
            or ""
        ).casefold()

        if apollo_status == "matched":
            apollo_cell.fill = PatternFill(
                fill_type="solid",
                fgColor="C6EFCE",
            )
        elif apollo_status in {
            "api_error",
            "connection_error",
        }:
            apollo_cell.fill = PatternFill(
                fill_type="solid",
                fgColor="FFC7CE",
            )

        # With one official per row, keep rows compact enough to scan.
        # The full Job Description remains available through wrapped text.
        description = str(
            job.job_description
            or ""
        )

        description_lines = max(
            1,
            description.count("\n") + 1,
        )

        worksheet.row_dimensions[
            row_index
        ].height = min(
            120,
            max(
                42,
                description_lines * 14,
            ),
        )

    def _format_sheet(
        self,
        worksheet: Any,
    ) -> None:
        worksheet.freeze_panes = "A2"

        last_column = get_column_letter(
            len(self.HEADERS)
        )

        worksheet.auto_filter.ref = (
            f"A1:{last_column}"
            f"{max(1, worksheet.max_row)}"
        )

        worksheet.sheet_view.showGridLines = False

        for index in range(
            1,
            len(self.HEADERS) + 1,
        ):
            letter = get_column_letter(index)
            worksheet.column_dimensions[
                letter
            ].width = 20

        custom_widths_by_header = {
            "Job ID": 14,
            "Job Title": 38,
            "Extracted Company Name": 28,
            "Location": 28,
            "Post Contact Name": 25,
            "Post Contact Email ID": 36,
            "Job Description": 70,
            "Apollo Company Name": 30,
            "Company Website": 35,
            "Industry": 28,
            "Company City": 22,
            "Company Country": 22,
            "Official Name": 28,
            "Official Title": 38,
            "Official Work Email": 34,
            "Official LinkedIn": 38,
            "Apollo Enrichment Error": 48,
            "Date": 14,
        }

        for header, width in (
            custom_widths_by_header.items()
        ):
            letter = get_column_letter(
                self._column_index(header)
            )
            worksheet.column_dimensions[
                letter
            ].width = width

        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0

    def _sync_table(
        self,
        worksheet: Any,
    ) -> None:
        last_data_row = worksheet.max_row
        table_name = "JobPostingsTable"

        if last_data_row < 2:
            if table_name in worksheet.tables:
                del worksheet.tables[table_name]
            return

        last_column = get_column_letter(
            len(self.HEADERS)
        )
        table_ref = (
            f"A1:{last_column}{last_data_row}"
        )

        if table_name in worksheet.tables:
            worksheet.tables[table_name].ref = (
                table_ref
            )
            return

        table = Table(
            displayName=table_name,
            ref=table_ref,
        )

        table.tableStyleInfo = (
            TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=False,
                showColumnStripes=False,
            )
        )

        worksheet.add_table(table)

    @classmethod
    def _column_index(
        cls,
        header: str,
    ) -> int:
        return cls.HEADERS.index(header) + 1

    @staticmethod
    def _thin_border() -> Border:
        side = Side(
            style="thin",
            color="D9E2F3",
        )

        return Border(
            left=side,
            right=side,
            top=side,
            bottom=side,
        )

    @staticmethod
    def _join_values(
        values: list[str]
        | tuple[str, ...]
        | None,
    ) -> str:
        if not values:
            return ""

        result: list[str] = []
        seen: set[str] = set()

        for value in values:
            cleaned = str(value).strip()

            if not cleaned:
                continue

            key = cleaned.casefold()

            if key not in seen:
                seen.add(key)
                result.append(cleaned)

        return ", ".join(result)

    @staticmethod
    def _safe_excel_value(
        value: Any,
    ) -> Any:
        if value is None:
            return ""

        if isinstance(value, str):
            cleaned = value.strip()

            if cleaned.startswith(
                (
                    "=",
                    "+",
                    "-",
                    "@",
                )
            ):
                return "'" + cleaned

            return cleaned

        return value